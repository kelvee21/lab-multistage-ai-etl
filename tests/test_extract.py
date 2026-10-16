"""reader tests: day tabs, typo tab names, quarantines."""
from __future__ import annotations

import openpyxl
import pytest

from etl.extract import ExtractionError, extract_all
from tests.helpers import day_only_contract, make_day_workbook

ITEMS = [
    ("MULTIGRAIN LOAF", {"K'KIRO": 6, "RBERO": 8, "N'RAMA": 25, "KGU": 3, "G YEAR": 0}),
    ("WHITE LOAF", {"CHIC": None, "FZONE": 1}),
]


def test_day_rows_keep_only_recorded_cells(tmp_path):
    contract = day_only_contract()
    make_day_workbook(tmp_path, "PROD.xlsx", {"3RD JUNE": ITEMS})
    result = extract_all(contract, tmp_path)

    assert len(result.rows) == 6  # 5 cells on item 1, 1 on item 2
    zero_row = next(r for r in result.rows if r["branch"] == "G YEAR")
    assert zero_row["qty"] == 0.0  # a typed 0 is data, we keep it
    assert all(r["date"] == "2026-06-03" for r in result.rows)
    assert all(r["fact"] == "production" for r in result.rows)
    assert all(r["source_sheet"] == "3RD JUNE" for r in result.rows)


def test_typo_sheet_letter_o_resolved(tmp_path):
    contract = day_only_contract()
    make_day_workbook(
        tmp_path, "PROD.xlsx", {"9TH JUNE": ITEMS, "1OTH JUNE": ITEMS[0:1]}
    )
    result = extract_all(contract, tmp_path)

    dates = {r["date"] for r in result.rows}
    assert dates == {"2026-06-09", "2026-06-10"}
    fixed = [n for n in result.notes if "letter-O" in n["message"]]
    assert len(fixed) == 1


def test_unknown_and_ignored_sheets(tmp_path):
    contract = day_only_contract()
    make_day_workbook(
        tmp_path, "PROD.xlsx", {"2ND JUNE": ITEMS, "WHATEVER": ITEMS, "SUMMARY": []}
    )
    result = extract_all(contract, tmp_path)

    kinds = result.sheets["production"]
    assert kinds["WHATEVER"] == "unknown_sheet"
    assert kinds["SUMMARY"] == "ignored_by_contract"
    assert kinds["2ND JUNE"].startswith("day_sheet")
    assert any(q["reason"] == "unknown_sheet" and q["source_sheet"] == "WHATEVER"
               for q in result.quarantines)


def test_stray_value_above_header_still_parses(tmp_path):
    contract = day_only_contract()
    make_day_workbook(tmp_path, "PROD.xlsx", {"22ND": ITEMS}, stray_value_above=True)
    result = extract_all(contract, tmp_path)

    assert result.rows
    assert all(r["date"] == "2026-06-22" for r in result.rows)
    assert not any(q["reason"] == "missing_product" for q in result.quarantines)


def test_missing_product_quarantined(tmp_path):
    contract = day_only_contract()
    make_day_workbook(tmp_path, "PROD.xlsx", {"2ND JUNE": []})
    # a branch value with no item name, same shape as the real tab
    path = tmp_path / "PROD.xlsx"
    wb = openpyxl.load_workbook(path)
    ws = wb["2ND JUNE"]
    ws.cell(row=2, column=3, value=7.0)
    wb.save(path)

    result = extract_all(contract, tmp_path)
    assert any(q["reason"] == "missing_product" for q in result.quarantines)
    assert result.rows == []


def test_negative_and_unparsable_numbers(tmp_path):
    contract = day_only_contract()
    make_day_workbook(
        tmp_path, "PROD.xlsx",
        {"2ND JUNE": [("MULTIGRAIN LOAF", {"K'KIRO": -2, "RBERO": "12,5"})]},
    )
    result = extract_all(contract, tmp_path)

    reasons = {q["reason"] for q in result.quarantines}
    assert "negative_qty" in reasons and "non_numeric_qty" in reasons
    assert result.rows == []


def test_thousands_separator_string(tmp_path):
    contract = day_only_contract()
    make_day_workbook(
        tmp_path, "PROD.xlsx",
        {"2ND JUNE": [("MULTIGRAIN LOAF", {"K'KIRO": "247,200"})]},
    )
    result = extract_all(contract, tmp_path)
    assert result.rows[0]["qty"] == 247200.0


def test_blank_sheet_reported_not_loaded(tmp_path):
    contract = day_only_contract()
    make_day_workbook(tmp_path, "PROD.xlsx", {"2ND JUNE": "empty", "1ST JUNE": ITEMS})
    result = extract_all(contract, tmp_path)

    kinds = result.sheets["production"]
    assert kinds["2ND JUNE"] == "day_sheet_no_recorded_values"
    assert {r["date"] for r in result.rows} == {"2026-06-01"}


def test_unmapped_column_values_quarantined(tmp_path):
    contract = day_only_contract()
    make_day_workbook(
        tmp_path, "PROD.xlsx",
        {"2ND JUNE": [
            ("MULTIGRAIN LOAF", {"K'KIRO": 3, "MYSTERY": 9}),
            ("WHITE LOAF", {"CHIC": 1}),
        ]},
        extra_header="MYSTERY",
    )
    result = extract_all(contract, tmp_path)

    assert any(q["reason"] == "unmapped_column" and q.get("product") == "MULTIGRAIN LOAF"
               for q in result.quarantines)
    white_loaf = [r for r in result.rows if r["product"] == "WHITE LOAF"]
    assert white_loaf and white_loaf[0]["branch"] == "CHIC"


def test_missing_source_file_fails_fast(tmp_path):
    contract = day_only_contract()
    with pytest.raises(ExtractionError):
        extract_all(contract, tmp_path)


def test_ignored_derived_column_note(tmp_path):
    contract = day_only_contract()
    make_day_workbook(
        tmp_path, "PROD.xlsx",
        {"2ND JUNE": [("MULTIGRAIN LOAF", {"K'KIRO": 3, "CHIC": 1, "TOTAL": 4})]},
        extra_header="TOTAL",
    )
    result = extract_all(contract, tmp_path)

    # the derived total is out by contract, so we only note it
    assert result.rows and all(r["branch"] in {"K'KIRO", "CHIC"} for r in result.rows)
    assert any("ignored derived column 'TOTAL'" in n["message"] and n["level"] == "info"
               for n in result.notes)
    assert not any(q["reason"] == "unmapped_column" for q in result.quarantines)


def test_row_total_mismatch_note(tmp_path):
    contract = day_only_contract()
    make_day_workbook(
        tmp_path, "PROD.xlsx",
        {"2ND JUNE": [("MULTIGRAIN LOAF", {"K'KIRO": 3, "CHIC": 1, "TOTAL": 30})]},
        extra_header="TOTAL",
    )
    result = extract_all(contract, tmp_path)

    mismatch_notes = [n for n in result.notes if "row(s) where the sheet's own" in n["message"]]
    assert len(mismatch_notes) == 1 and mismatch_notes[0]["level"] == "anomaly"
    assert "MULTIGRAIN LOAF" in mismatch_notes[0]["message"]
    # the row still loads, the sheet's total is a hint not a gate
    assert len([r for r in result.rows if r["product"] == "MULTIGRAIN LOAF"]) == 2


def test_aggregate_total_row_never_a_product(tmp_path):
    contract = day_only_contract()
    make_day_workbook(
        tmp_path, "PROD.xlsx",
        {"2ND JUNE": [
            ("MULTIGRAIN LOAF", {"K'KIRO": 6}),
            ("TOTAL", {"K'KIRO": 99}),      # stray value inside the aggregate row
            ("TOTAL", {}),                  # normal aggregate row: no values
        ]},
    )
    result = extract_all(contract, tmp_path)

    assert [r["product"] for r in result.rows] == ["MULTIGRAIN LOAF"]
    strays = [q for q in result.quarantines if q["reason"] == "aggregate_row_value"]
    assert len(strays) == 1 and strays[0]["product"] == "TOTAL" and strays[0]["branch"] == "K'KIRO"


def test_branch_alias_loads_and_notes(tmp_path):
    contract = day_only_contract()
    contract["sources"][0]["branch_aliases"] = {"GDYEAR": "G YEAR"}
    make_day_workbook(
        tmp_path, "PROD.xlsx",
        {"2ND JUNE": [("MULTIGRAIN LOAF", {"G YEAR": 5})]},
    )
    # retype the G YEAR header the way the real waste 14th tab does
    from tests.helpers import BRANCHES
    path = tmp_path / "PROD.xlsx"
    wb = openpyxl.load_workbook(path)
    ws = wb["2ND JUNE"]
    gy_col = 3 + BRANCHES.index("G YEAR")
    ws.cell(row=1, column=gy_col, value="GDYEAR")
    wb.save(path)

    result = extract_all(contract, tmp_path)
    gy_rows = [r for r in result.rows if r["branch"] == "G YEAR"]
    assert gy_rows and gy_rows[0]["qty"] == 5  # loaded under the correct name
    assert any("GDYEAR -> G YEAR" in n["message"] for n in result.notes)


def test_unnamed_column_with_values_noted(tmp_path):
    contract = day_only_contract()
    make_day_workbook(
        tmp_path, "PROD.xlsx",
        {"2ND JUNE": [("MULTIGRAIN LOAF", {"K'KIRO": 2})]},
    )
    # a value under an empty header. we cannot name it so we only note it.
    path = tmp_path / "PROD.xlsx"
    wb = openpyxl.load_workbook(path)
    ws = wb["2ND JUNE"]
    ws.cell(row=2, column=3 + len(contract["branch_headers"]) + 1, value=77.0)
    wb.save(path)

    result = extract_all(contract, tmp_path)
    unnamed_notes = [n for n in result.notes if "unnamed column" in n["message"]]
    assert len(unnamed_notes) == 1 and unnamed_notes[0]["level"] == "anomaly"
    assert not any(q["reason"] == "unmapped_column" for q in result.quarantines)