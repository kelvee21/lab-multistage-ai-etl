"""Synthetic workbook builders for the tests.

these are fake on purpose. the real bakery files stay in data/raw and are never
committed, so the tests have to build their own small workbooks.
"""
from __future__ import annotations

import openpyxl

BRANCHES = ["HQ", "K'KIRO", "RBERO", "NBOYE", "CHIC", "M'HIMA", "N'RAMA", "KGU",
            "RMRA", "K'IMENTI", "G YEAR", "FZONE", "VISTA", "C SHOP"]


def day_sheet(ws, item_rows: list[tuple[str, dict]], header_row: int = 1,
              stray_value_above: bool = False, extra_header: str | None = None) -> None:
    """One day tab: col A row numbers, col B ITEM, branches from col C.

    stray_value_above copies the real dispatch 22nd tab, a number sitting in
    column A above the header.
    extra_header adds one column after the branches. put the value in
    values[extra_header] and it lands under that header.
    """
    hr = header_row + 1 if stray_value_above else header_row
    if stray_value_above:
        ws.cell(row=1, column=1, value=4.0)
    ws.cell(row=hr, column=1, value=None)
    ws.cell(row=hr, column=2, value="ITEM")
    for i, branch in enumerate(BRANCHES):
        ws.cell(row=hr, column=3 + i, value=branch)
    if extra_header is not None:
        ws.cell(row=hr, column=3 + len(BRANCHES), value=extra_header)

    for j, (product, values) in enumerate(item_rows, start=1):
        row = hr + j
        ws.cell(row=row, column=1, value=j)
        ws.cell(row=row, column=2, value=product)
        for i, branch in enumerate(BRANCHES):
            if branch in values:
                value = values[branch]
                if value is not None:
                    ws.cell(row=row, column=3 + i, value=value)
        if extra_header is not None and extra_header in values:
            ws.cell(row=row, column=3 + len(BRANCHES), value=values[extra_header])


def make_day_workbook(tmp_path, filename: str, sheets: dict[str, list[tuple]], **kwargs) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for tab, item_rows in sheets.items():
        ws = wb.create_sheet(title=tab)
        if item_rows == "empty":
            day_sheet(ws, [], **kwargs)
        else:
            day_sheet(ws, item_rows, **kwargs)
    wb.save(tmp_path / filename)


def make_sales_workbook(tmp_path, filename: str, header: list[str], rows: list[list]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales Report"
    ws.append(header)
    for row in rows:
        ws.append(row)
    pivot = wb.create_sheet("Sheet1")
    pivot.cell(row=1, column=1, value="Row Labels")  # deliberately not transaction-shaped
    wb.save(tmp_path / filename)


SALES_HEADER_SEP = [
    "Document No.", "Date", "Item", "Partner", "Group", "Location", "Quantity",
    "Measure", "Sale price", "VAT included", "Sale amount", "Note", "Operation",
]


def tx_row(doc=200394776, date="22.09.2023", item="BS MULTIGRAIN LOAF", partner="RETAIL",
           location="SAWA CITI HQ", qty=1, price=4000, amount=4000, operation="Sale"):
    return [doc, date, item, partner, "Default group", location, qty, "item",
            price, round(price * 5 / 23, 12), amount, None, operation]


def base_contract():
    """same shape as the real contract, but pointing at the fake file names."""
    return {
        "version": "test-1",
        "branch_headers": BRANCHES,
        "sources": [
            {
                "id": "production",
                "file": "PROD.xlsx",
                "fact": "production",
                "grid": "day",
                "period": {"year": 2026, "month": 6},
                "expected_days": {"from": 1, "to": 3},
                "day_sheet": {"regex": "^(\\d{1,2})(?:ST|ND|RD|TH)?\\s*(?:JUN|JUNE)?\\s*$", "fix_letter_O": True},
                "header_search_rows": 3,
                "item_header": "ITEM",
                "branch_aliases": {},
                "ignore_columns": {"TOTAL": "per-row total (derived)", "PRICE": "unit price"},
                "row_total_column": "TOTAL",
                "aggregate_rows": ["TOTAL"],
                "ignore_sheets": {"SUMMARY": "currency totals, excluded by contract"},
                "notes": [],
            },
            {
                "id": "sales_sep",
                "file": "SALES.xlsx",
                "fact": "sales",
                "grid": "transactions",
                "sheet": "Sales Report",
                "ignore_sheets": {"Sheet1": "pivot"},
                "header_search_rows": 3,
                "columns": {"date": "Date", "product": "Item", "branch": "Location",
                            "qty": "Quantity", "amount": "Sale amount", "partner": "Partner",
                            "operation": "Operation"},
                "date_format": "%d.%m.%Y",
                "operation_accept": ["Sale"],
                "notes": [],
            },
        ],
    }


def day_only_contract():
    contract = base_contract()
    contract["sources"] = [contract["sources"][0]]
    return contract


def sales_only_contract():
    contract = base_contract()
    contract["sources"] = [contract["sources"][1]]
    return contract
